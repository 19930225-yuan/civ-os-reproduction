# -*- coding: utf-8 -*-
"""第二轮修改（遗留问题）+ 生成接受全部修订的清洁终稿"""
import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = "AI编辑"; DATE = "2026-10-17T04:30:00Z"
_c = [2000]
def nid():
    _c[0] += 1
    return str(_c[0])

def _set_run_text(run_el, text):
    for ch in list(run_el):
        if ch.tag in (qn('w:t'), qn('w:delText')):
            run_el.remove(ch)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    run_el.append(t)

def _wrap_del(run_el):
    d = OxmlElement('w:del')
    d.set(qn('w:id'), nid()); d.set(qn('w:author'), AUTHOR); d.set(qn('w:date'), DATE)
    for t in run_el.findall(qn('w:t')):
        t.tag = qn('w:delText')
    d.append(run_el); return d

def _wrap_ins(run_el, text):
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    _set_run_text(run_el, text); ins.append(run_el); return ins

def replace_in_para(p, old, new):
    for run in p.runs:
        if old in run.text:
            before, after = run.text.split(old, 1)
            r = run._r
            r_old = copy.deepcopy(r); _set_run_text(r_old, old)
            r_new = copy.deepcopy(r); _set_run_text(r_new, new)
            r_after = copy.deepcopy(r); _set_run_text(r_after, after)
            _set_run_text(r, before)
            parent = r.getparent(); idx = list(parent).index(r)
            nodes = ([r] if before else []) + [_wrap_del(r_old), _wrap_ins(r_new, new)] + ([r_after] if after else [])
            for off, n in enumerate(nodes):
                parent.insert(idx + off, n)
            if not before:
                parent.remove(r)
            return True
    return False

def replace_para_tracked(p, new):
    runs = [r._r for r in p.runs]
    if not runs:
        return
    first = runs[0]; parent = first.getparent(); idx = list(parent).index(first)
    for r in runs:
        parent.remove(r)
    d = OxmlElement('w:del')
    d.set(qn('w:id'), nid()); d.set(qn('w:author'), AUTHOR); d.set(qn('w:date'), DATE)
    for r in runs:
        for t in r.findall(qn('w:t')):
            t.tag = qn('w:delText')
        d.append(r)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    rn = copy.deepcopy(first); _set_run_text(rn, new); ins.append(rn)
    parent.insert(idx, d); parent.insert(idx + 1, ins)

def insert_para_tracked(doc, ref_p, text, style_name, after=True):
    new_p = doc.add_paragraph(style=style_name)
    r = new_p.add_run(text)._r
    parent = r.getparent(); idx = list(parent).index(r)
    parent.remove(r)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    ins.append(r); parent.insert(idx, ins)
    pPr = new_p._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        st = pPr.find(qn('w:pStyle'))
        (st.addnext(rPr) if st is not None else pPr.insert(0, rPr))
    m = OxmlElement('w:ins')
    m.set(qn('w:id'), nid()); m.set(qn('w:author'), AUTHOR); m.set(qn('w:date'), DATE)
    rPr.append(m)
    (ref_p._p.addnext(new_p._p) if after else ref_p._p.addprevious(new_p._p))
    return new_p

def find_para(doc, needle):
    hits = [p for p in doc.paragraphs if needle in p.text]
    assert len(hits) == 1, f"{needle!r} -> {len(hits)}"
    return hits[0]

def tr(p, old, new, tag, log):
    if not replace_in_para(p, old, new):
        replace_para_tracked(p, p.text.replace(old, new))
    log.append(tag)

# ================= 正文第二轮 =================
d = Document("/home/work/.openclaw/workspace/投稿69_修改稿.docx")
log = []
P = d.paragraphs

# N1 印度案例判定理由重写
india = [p for p in P if p.text.startswith("印度（外壳=3.5）")]
assert len(india) == 1
replace_para_tracked(india[0],
    "印度（外壳=3.5）\xa0：脆弱性未转化为断裂。按上述三项判定指标，指标（b）冲击者未建立替代性核心认同"
    "与指标（c）冲击者未系统性清除核心经典与仪式均告成立——核心经典（吠陀、奥义书）未被替代，"
    "种姓网络未被伊斯兰教法体系替代，梵文未被阿拉伯文或波斯文替代——两项成立即属“非极端冲击”。"
    "需要说明的是，伊斯兰政权在印度次大陆的统治延续数百年，指标（a）“冲击持续时间不足三代人”并不成立；"
    "因此印度的“未穿透”应归因于冲击者未建立替代性核心认同、未系统性清除核心经典与仪式，而非冲击时间短——"
    "判定“非极端”的关键在于替代的深度与系统性，而非时间长短。印度三项操作化指标的逐条论证详见补充材料C第C.16.1节。")
log.append("N1 印度案例改以指标(b)(c)立论，删除对“冲击时间短”的依赖")

# N2 表11 F1/F2 通俗化 + 补表注
t11 = d.tables[10]
assert t11.rows[1].cells[0].text.strip() == "F1"
def cell_edit(cell, old, new):
    p = cell.paragraphs[0]
    if not replace_in_para(p, old, new):
        replace_para_tracked(p, new)
cell_edit(t11.rows[1].cells[1], t11.rows[1].cells[1].text,
          "出现“政权崩溃后未能恢复独立宏观政治组织”但五条件中仍有≥1个条件得分≥4.0的案例")
cell_edit(t11.rows[1].cells[2], t11.rows[1].cells[2].text,
          "未触发（拜占庭、古埃及得分≥4的条件数均为0）")
cell_edit(t11.rows[2].cells[1], t11.rows[2].cells[1].text,
          "出现外壳调整得分＞1.0的补充消亡案例（即高外壳同样走向消亡）")
log.append("N2 表11 F1/F2表述通俗化")
tbl_el = t11._tbl
nxt = tbl_el.getnext()
ref_p = None
from docx.text.paragraph import Paragraph
if nxt is not None and nxt.tag == qn('w:p'):
    ref_p = Paragraph(nxt, d)
note = insert_para_tracked(d, ref_p, "注：F1检验“条件强度与政治架构恢复”之间的关联，F2检验“外壳调整与消亡”的对应关系；"
                          "两者均为单侧检验——框架不声称高外壳足以保证存续，亦不声称外壳是唯一决定因素。",
                          "表注", after=False)
log.append("N2b 表11补表注（说明F1/F2的检验对象）")

# N3 预测四补适用条件
tr(P[212], "检验核心功能件脱离母体后的独立存续能力。",
   "检验核心功能件脱离母体后的独立存续能力。该预测仅在移民族群脱离母体制度、无宗族或教会等母体组织依托的散居条件下成立。",
   "N3 预测四补适用条件", log)

# N4 §8.1 编码信度补可复核性说明
tr(P[172], "编码信度建立在“推理路径的透明化”而非“评分者间的一致性”之上。",
   "编码信度建立在“推理路径的透明化”而非“评分者间的一致性”之上。全部评分规则、逐条判定理由与反向证据均已公开于补充材料A、C，可供同行独立复核与再编码。",
   "N4 §8.1补可复核性说明", log)

d.save("/home/work/.openclaw/workspace/投稿69_修改稿.docx")
print("正文第二轮："); [print(" -", x) for x in log]

# ================= 补充材料第二轮 =================
d2 = Document("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")
log2 = []
tr(find_para(d2, "（c）梵文未被阿拉伯文或波斯文替代"),
   "但梵文在婆罗门精英中的传承未中断。",
   "但梵文在婆罗门精英中的传承未中断。上述三条证据分别对应判定规则中的指标（b）与（c）；"
   "伊斯兰政权在印度次大陆的统治延续数百年，指标（a）并不成立，故本文不以“冲击时间短”立论，"
   "而以“替代的深度与系统性不足”立论。",
   "T1 C.16.1同步修正印度判定理由", log2)
d2.save("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")
print("补充材料第二轮："); [print(" -", x) for x in log2]

# ================= 接受全部修订 -> 清洁终稿 =================
def accept_all(src, dst):
    dd = Document(src)
    body = dd.element.body
    for ins in list(body.findall('.//' + qn('w:ins'))):
        parent = ins.getparent()
        idx = list(parent).index(ins)
        if parent.tag == qn('w:rPr'):      # 段落标记插入 -> 仅移除标记
            parent.remove(ins)
            continue
        for ch in list(ins):
            parent.insert(idx, ch); idx += 1
        parent.remove(ins)
    for dl in list(body.findall('.//' + qn('w:del'))):
        parent = dl.getparent()
        if parent.tag == qn('w:rPr'):
            # 段落标记删除：连同空段落一起移除
            p = parent.getparent().getparent()
            has_text = any(t.text for t in p.findall('.//' + qn('w:t')))
            if not has_text:
                gp = p.getparent(); gp.remove(p)
            else:
                parent.remove(dl)
            continue
        parent.remove(dl)
    dd.save(dst)
    return dd

accept_all("/home/work/.openclaw/workspace/投稿69_修改稿.docx", "/home/work/.openclaw/workspace/投稿69_终稿.docx")
accept_all("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx", "/home/work/.openclaw/workspace/网络补充材料_终稿.docx")
print("clean finals saved")

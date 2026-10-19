# -*- coding: utf-8 -*-
"""复核（含修订态文本）+ 重出终稿"""
from docx import Document
from docx.oxml.ns import qn

def cell_text(cell):
    out = []
    for el in cell._tc.iter():
        tag = el.tag.split('}')[1]
        if tag == 't':
            anc = el.getparent().getparent()
            if anc is not None and anc.tag == qn('w:del'):
                continue
            out.append(el.text or '')
    return ''.join(out)

def num(s):
    s = s.strip()
    return float(s) if s else None

sup = Document("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")
T = sup.tables
ok_all = True
print("== 复核（读取修订态最终文本）==")
for ti, name, totrow in [(7, "条件一", 6), (10, "条件二", 6)]:
    t = T[ti]
    for ci in range(1, 9):
        case = cell_text(t.rows[0].cells[ci]).strip()
        vals = [num(cell_text(t.rows[r].cells[ci])) for r in range(1, 6)]
        s = sum(vals)
        tot = num(cell_text(t.rows[totrow].cells[ci]))
        flag = "✓" if abs(s - tot) < 1e-9 else "✗"
        if flag == "✗": ok_all = False
        print(f"  {flag} {name} {case}: {'+'.join(str(v) for v in vals)} = {s:.1f} vs 总分{tot}")
t = T[25]
for ri in range(1, 9):
    case = cell_text(t.rows[ri].cells[0]).strip()
    mi = [num(x) for x in cell_text(t.rows[ri].cells[1]).split("/")]
    ma = [num(x) for x in cell_text(t.rows[ri].cells[3]).split("/")]
    mtt = num(cell_text(t.rows[ri].cells[2])); mat = num(cell_text(t.rows[ri].cells[4]))
    org = num(cell_text(t.rows[ri].cells[5]))
    ok = (len(mi) == 5 and len(ma) == 5 and abs(sum(mi) - mtt) < 1e-9
          and abs(sum(ma) - mat) < 1e-9 and abs(mtt + mat - org) < 1e-9)
    ok_all = ok_all and ok
    print(f"  {'✓' if ok else '✗'} 条件四 {case}: 微{sum(mi):.1f}/{mtt} 宏{sum(ma):.1f}/{mat} 总{mtt+mat:.1f}/{org}")
# 苏美尔
b1 = T[39]
dep = cell_text(b1.rows[1].cells[2])
import re
vals = [float(x) for x in re.findall(r"=\s*([0-9.]+)", dep)]
score = num(cell_text(b1.rows[1].cells[1]))
ok = abs(sum(vals) - score) < 1e-9
ok_all = ok_all and ok
print(f"  {'✓' if ok else '✗'} 苏美尔信仰: {vals} 和{sum(vals):.1f} vs 公布{score}")
# ≥4数
t = T[36]
for ri in range(1, 9):
    case = cell_text(t.rows[ri].cells[0]).strip()
    scores = [num(cell_text(t.rows[ri].cells[ci])) for ci in range(1, 6)]
    ge4 = sum(1 for v in scores if v >= 4)
    pub = num(cell_text(t.rows[ri].cells[6]))
    ok = ge4 == pub
    ok_all = ok_all and ok
    print(f"  {'✓' if ok else '✗'} ≥4数 {case}: 实算{ge4} vs 表载{pub}")
print("  >>> 全部一致 ✓" if ok_all else "  >>> 仍有不一致 ✗")

main = Document("/home/work/.openclaw/workspace/投稿69_修改稿.docx")
t6 = main.tables[5]
hdr = [cell_text(c).strip() for c in t6.rows[0].cells]
col = hdr.index("≥4数")
for ri in range(1, len(t6.rows)):
    case = cell_text(t6.rows[ri].cells[0]).strip()
    scores = [num(cell_text(t6.rows[ri].cells[ci])) for ci in range(1, 6)]
    ge4 = sum(1 for v in scores if v >= 4)
    pub = num(cell_text(t6.rows[ri].cells[col]))
    print(f"  {'✓' if ge4 == pub else '✗'} 主文表6 ≥4数 {case}: {ge4} vs {pub}")

def accept_all(src, dst):
    dd = Document(src)
    body = dd.element.body
    for ins in list(body.findall('.//' + qn('w:ins'))):
        parent = ins.getparent(); idx = list(parent).index(ins)
        if parent.tag == qn('w:rPr'):
            parent.remove(ins); continue
        for ch in list(ins):
            parent.insert(idx, ch); idx += 1
        parent.remove(ins)
    for dl in list(body.findall('.//' + qn('w:del'))):
        parent = dl.getparent()
        if parent.tag == qn('w:rPr'):
            p = parent.getparent().getparent()
            has_text = any(t.text for t in p.findall('.//' + qn('w:t')))
            if not has_text:
                p.getparent().remove(p)
            else:
                parent.remove(dl)
            continue
        parent.remove(dl)
    dd.save(dst)

accept_all("/home/work/.openclaw/workspace/投稿69_修改稿.docx", "/home/work/.openclaw/workspace/投稿69_终稿.docx")
accept_all("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx", "/home/work/.openclaw/workspace/网络补充材料_终稿.docx")
print("终稿已更新")

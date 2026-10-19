# -*- coding: utf-8 -*-
"""第三轮：印度案例补入巴克提运动/锡克教证据 + 重出终稿"""
import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = "AI编辑"; DATE = "2026-10-17T05:00:00Z"
_c = [3000]
def nid():
    _c[0] += 1
    return str(_c[0])

def _set_run_text(run_el, text):
    for ch in list(run_el):
        if ch.tag in (qn('w:t'), qn('w:delText')):
            run_el.remove(ch)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    run_el.append(t)

def _wrap_del(run_el):
    d = OxmlElement('w:del')
    d.set(qn('w:id'), nid()); d.set(qn('w:author'), AUTHOR); d.set(qn('w:date'), DATE)
    for t in run_el.findall(qn('w:t')):
        t.tag = qn('w:delText')
    d.append(run_el); return d

def _wrap_ins(run_el, text):
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    _set_run_text(run_el, text); ins.append(run_el); return ins

def replace_in_para(p, old, new):
    for run in p.runs:
        if old in run.text:
            before, after = run.text.split(old, 1)
            r = run._r
            r_old = copy.deepcopy(r); _set_run_text(r_old, old)
            r_new = copy.deepcopy(r); _set_run_text(r_new, new)
            r_after = copy.deepcopy(r); _set_run_text(r_after, after)
            _set_run_text(r, before)
            parent = r.getparent(); idx = list(parent).index(r)
            nodes = ([r] if before else []) + [_wrap_del(r_old), _wrap_ins(r_new, new)] + ([r_after] if after else [])
            for off, n in enumerate(nodes):
                parent.insert(idx + off, n)
            if not before:
                parent.remove(r)
            return True
    return False

def insert_para_tracked(doc, ref_p, text, style_name, after=True):
    new_p = doc.add_paragraph(style=style_name)
    r = new_p.add_run(text)._r
    parent = r.getparent(); idx = list(parent).index(r)
    parent.remove(r)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    ins.append(r); parent.insert(idx, ins)
    pPr = new_p._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        st = pPr.find(qn('w:pStyle'))
        (st.addnext(rPr) if st is not None else pPr.insert(0, rPr))
    m = OxmlElement('w:ins')
    m.set(qn('w:id'), nid()); m.set(qn('w:author'), AUTHOR); m.set(qn('w:date'), DATE)
    rPr.append(m)
    (ref_p._p.addnext(new_p._p) if after else ref_p._p.addprevious(new_p._p))
    return new_p

def find_para(doc, needle):
    hits = [p for p in doc.paragraphs if needle in p.text]
    assert len(hits) == 1, f"{needle!r} -> {len(hits)}"
    return hits[0]

def flatten_para(p):
    """把段落内已有的修订痕迹就地接受，使其可继续编辑"""
    for ins in list(p._p.findall(qn('w:ins'))):
        parent = ins.getparent(); idx = list(parent).index(ins)
        if parent.tag == qn('w:rPr'):
            parent.remove(ins); continue
        for ch in list(ins):
            parent.insert(idx, ch); idx += 1
        parent.remove(ins)
    for dl in list(p._p.findall(qn('w:del'))):
        parent = dl.getparent()
        if parent.tag == qn('w:rPr'):
            parent.remove(dl); continue
        parent.remove(dl)

def full_text(p):
    out = []
    for el in p._p:
        tag = el.tag.split('}')[1]
        if tag == 'r':
            out.append(''.join(t.text or '' for t in el.findall(qn('w:t'))))
        elif tag == 'ins':
            for r in el.findall(qn('w:r')):
                out.append(''.join(t.text or '' for t in r.findall(qn('w:t'))))
    return ''.join(out)

# ---------- 正文 ----------
d = Document("/home/work/.openclaw/workspace/投稿69_修改稿.docx")
p_india = [p for p in d.paragraphs if full_text(p).startswith("印度（外壳=3.5）")]
assert len(p_india) == 1
old = "印度三项操作化指标的逐条论证详见补充材料C第C.16.1节。"
new = ("更值得注意的是，印度文明对伊斯兰压力并非被动承受——中世纪巴克提运动以虔信改革回应宗教竞争；"
       "锡克教于15世纪末兴起于旁遮普，在莫卧儿统治下发展出宗教—军事合一的组织形态（戈宾德·辛格1699年建立卡尔沙），"
       "并由兰季特·辛格于18—19世纪之交在旁遮普重建独立宏观政治组织。这正是外壳调整“方向多样性”（5d=1.0）与"
       "“坚韧性”（5c=1.0）的具体体现，也与表4中印度“有条件重建”的判定相互印证。"
       "上述判定与证据的逐条论证详见补充材料C第C.16.1节。")
assert replace_in_para(flatten_para(p_india[0]) or p_india[0], old, new)
print("R1 正文§8.5补入巴克提运动/锡克教证据")

# ---------- 补充材料 ----------
d2 = Document("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")
pc = find_para(d2, "（c）梵文未被阿拉伯文或波斯文替代")
flatten_para(pc)
insert_para_tracked(d2, pc,
    "（d）压力下的宗教创新：巴克提运动与锡克教。印度文明面对伊斯兰化压力并非被动承受：中世纪巴克提运动以虔信改革"
    "回应宗教竞争；锡克教于15世纪末兴起于旁遮普，兼受印度教巴克提传统与伊斯兰苏菲主义影响，在莫卧儿统治下逐步强化"
    "自我防卫——戈宾德·辛格于1699年建立卡尔沙（Khalsa），确立宗教—军事合一的组织形态。18—19世纪之交，"
    "锡克诸邦在旁遮普重建了独立宏观政治组织（兰季特·辛格政权）。上述创新分别印证5c=1.0（调整的坚韧性）与"
    "5d=1.0（调整的方向多样性），并为表4中印度“微观自保+有条件重建”的判定提供案例支持；但其组织能量集中于旁遮普一隅，"
    "未形成跨区域的制度化重建通道（5e=0.5），不改变印度外壳总分3.5的判定。",
    "Normal", after=True)
print("R2 补充材料C.16.1补入(d)条证据")
d.save("/home/work/.openclaw/workspace/投稿69_修改稿.docx")
d2.save("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")

# ---------- 重出终稿 ----------
def accept_all(src, dst):
    dd = Document(src)
    body = dd.element.body
    for ins in list(body.findall('.//' + qn('w:ins'))):
        parent = ins.getparent(); idx = list(parent).index(ins)
        if parent.tag == qn('w:rPr'):
            parent.remove(ins); continue
        for ch in list(ins):
            parent.insert(idx, ch); idx += 1
        parent.remove(ins)
    for dl in list(body.findall('.//' + qn('w:del'))):
        parent = dl.getparent()
        if parent.tag == qn('w:rPr'):
            p = parent.getparent().getparent()
            has_text = any(t.text for t in p.findall('.//' + qn('w:t')))
            if not has_text:
                p.getparent().remove(p)
            else:
                parent.remove(dl)
            continue
        parent.remove(dl)
    dd.save(dst)

accept_all("/home/work/.openclaw/workspace/投稿69_修改稿.docx", "/home/work/.openclaw/workspace/投稿69_终稿.docx")
accept_all("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx", "/home/work/.openclaw/workspace/网络补充材料_终稿.docx")
print("终稿已更新")

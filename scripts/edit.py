# -*- coding: utf-8 -*-
"""Apply tracked-change edits to 投稿69.docx and 网络补充材料.docx"""
import copy, os, sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = "AI编辑"
DATE = "2026-10-17T02:30:00Z"
_counter = [1000]

def nid():
    _counter[0] += 1
    return str(_counter[0])

def _set_run_text(run_el, text):
    for child in list(run_el):
        if child.tag in (qn('w:t'), qn('w:delText')):
            run_el.remove(child)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run_el.append(t)

def _wrap_del(run_el):
    d = OxmlElement('w:del')
    d.set(qn('w:id'), nid()); d.set(qn('w:author'), AUTHOR); d.set(qn('w:date'), DATE)
    for t in run_el.findall(qn('w:t')):
        t.tag = qn('w:delText')
    d.append(run_el)
    return d

def _wrap_ins(run_el, text):
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    _set_run_text(run_el, text)
    ins.append(run_el)
    return ins

def replace_in_para(p, old, new):
    """Tracked substring replacement inside one run."""
    for run in p.runs:
        if old in run.text:
            before, after = run.text.split(old, 1)
            r = run._r
            r_old = copy.deepcopy(r); _set_run_text(r_old, old)
            r_new = copy.deepcopy(r); _set_run_text(r_new, new)
            r_after = copy.deepcopy(r); _set_run_text(r_after, after)
            _set_run_text(r, before)
            parent = r.getparent()
            idx = list(parent).index(r)
            nodes = []
            if before:
                nodes.append(r)
            nodes.append(_wrap_del(r_old))
            nodes.append(_wrap_ins(r_new, new))
            if after:
                nodes.append(r_after)
            for off, n in enumerate(nodes):
                parent.insert(idx + off, n)
            if not before:
                parent.remove(r)
            return True
    return False

def replace_whole_run_tracked(run, new):
    r = run._r
    r_new = copy.deepcopy(r)
    r_del = copy.deepcopy(r)
    parent = r.getparent()
    idx = list(parent).index(r)
    parent.insert(idx, _wrap_del(r_del))
    parent.insert(idx + 1, _wrap_ins(r_new, new))
    parent.remove(r)

def replace_para_tracked(p, new):
    """Tracked whole-paragraph text replacement (fallback)."""
    runs = [r._r for r in p.runs]
    if not runs:
        return
    first = runs[0]
    parent = first.getparent()
    idx = list(parent).index(first)
    for r in runs:
        parent.remove(r)
    d = OxmlElement('w:del')
    d.set(qn('w:id'), nid()); d.set(qn('w:author'), AUTHOR); d.set(qn('w:date'), DATE)
    for r in runs:
        for t in r.findall(qn('w:t')):
            t.tag = qn('w:delText')
        d.append(r)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    r_new = copy.deepcopy(first)
    _set_run_text(r_new, new)
    ins.append(r_new)
    parent.insert(idx, d)
    parent.insert(idx + 1, ins)

def mark_para_deleted(p):
    for run in list(p.runs):
        r = run._r
        parent = r.getparent()
        idx = list(parent).index(r)
        parent.remove(r)
        parent.insert(idx, _wrap_del(r))
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is not None:
            pStyle.addnext(rPr)
        else:
            pPr.insert(0, rPr)
    d = OxmlElement('w:del')
    d.set(qn('w:id'), nid()); d.set(qn('w:author'), AUTHOR); d.set(qn('w:date'), DATE)
    rPr.append(d)

def insert_para_tracked(doc, ref_p, text, style_name, after=True):
    new_p = doc.add_paragraph(style=style_name)
    run = new_p.add_run(text)
    r = run._r
    parent = r.getparent()
    idx = list(parent).index(r)
    parent.remove(r)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    ins.append(r)
    parent.insert(idx, ins)
    pPr = new_p._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is not None:
            pStyle.addnext(rPr)
        else:
            pPr.insert(0, rPr)
    ins_mark = OxmlElement('w:ins')
    ins_mark.set(qn('w:id'), nid()); ins_mark.set(qn('w:author'), AUTHOR); ins_mark.set(qn('w:date'), DATE)
    rPr.append(ins_mark)
    if after:
        ref_p._p.addnext(new_p._p)
    else:
        ref_p._p.addprevious(new_p._p)
    return new_p

def find_para(doc, needle, style=None):
    hits = [p for p in doc.paragraphs if needle in p.text and (style is None or p.style.name == style)]
    assert len(hits) == 1, f"needle {needle!r} hits={len(hits)}"
    return hits[0]

# ---------------- 主文稿 ----------------
src = "投稿69.docx"; dst = "/home/work/.openclaw/workspace/投稿69_修改稿.docx"
d = Document(src)
P = d.paragraphs

log = []
def tr(p, old, new, tag):
    ok = replace_in_para(p, old, new)
    if not ok:
        replace_para_tracked(p, p.text.replace(old, new))
    log.append(tag)

# M1 摘要精简
body_run = [r for r in P[1].runs if "古埃及宏观组织能力高于犹太文明" in r.text]
assert len(body_run) == 1
new_abs = ("古埃及的宏观组织能力高于犹太文明，却在政治崩溃后彻底消亡；犹太文明在失去国土两千年后仍维持认同连续性。"
           "既有理论——曼的IEMP模型、文化记忆理论、路径依赖与韧性研究——均不区分“奠基期写入的结构性参数”与“后续行为性调整”，"
           "也不处理“核心承诺能否脱离政治载体”这一前置条件。本文提出“文明操作系统”（Civ-OS）框架，"
           "将文明视为由信仰锚点、血缘身份传递、文化认同体系、组织能力、外壳调整五个功能件构成的结构化应答系统，"
           "并采用“双轨锚定法”：结构性参数（5a/5b）锚定于奠基期制度文本，行为性参数（5c/5d/5e）锚定于全历史周期，"
           "以阻断“评分—结论”的循环风险。八个核心案例显示：外壳调整评分≤1.0的文明均未能在政权崩溃后恢复独立宏观政治组织；"
           "结构性参数完整的文明（犹太、中华）则走向“壳碎核移”。该框架与历史唯物主义构成分层互补关系。")
replace_whole_run_tracked(body_run[0], new_abs); log.append("M1 摘要重写")

# M2 英文摘要：整段重写 + 段落合并（[4] 段落跟踪删除）
en_run = [r for r in P[3].runs if "Ancient Egypt possessed" in r.text]
assert len(en_run) == 1
new_en = ("Ancient Egypt possessed higher macro-level organizational capacity than Jewish civilization, "
         "yet collapsed entirely after political disintegration, whereas the Jewish people maintained "
         "civilizational identity continuity for two millennia without a homeland. Existing theories—Mann's "
         "IEMP model, cultural memory theory, path dependency, and resilience studies—fail to distinguish "
         "between structural parameters written in at the founding stage and subsequent behavioral adjustments, "
         "nor do they address the prior condition of whether core commitments can be detached from political "
         "carriers. This article proposes the \u201cCivilization Operating System\u201d (Civ-OS) framework, "
         "conceptualizing civilization as a structured response system composed of five functional components—"
         "belief anchoring, bloodline identity transmission, cultural identity systems, organizational capacity, "
         "and shell adjustment—and adopts a dual-track anchoring method: structural parameters (5a/5b) are "
         "anchored in founding-period institutional texts, whereas behavioral parameters (5c/5d/5e) are anchored "
         "across the full historical cycle, thereby methodologically blocking the circular risk of "
         "\u201cscoring–conclusion\u201d loops. Eight core cases demonstrate that civilizations with "
         "shell-adjustment scores of 1.0 or lower all failed to restore independent macro-political organization "
         "after regime collapse, whereas those with intact structural parameters—Jewish and Chinese civilization—"
         "moved toward \u201cshell fracture with core migration.\u201d This framework constitutes a layered "
         "complementarity with historical materialism.")
replace_whole_run_tracked(en_run[0], new_en)
P[3].style = d.styles['英文摘要']
mark_para_deleted(P[4])
P[5].style = d.styles['英文关键词']
log.append("M2/M3 英文摘要合并重写、样式修正")

tr(P[8], "依然保持同一文明的名义维持认同连续性", "依然以同一文明的名义维持着认同连续性", "M4 病句修正")
tr(P[41], "第五节以三个消亡案例和一个极端存续案例验证外壳断裂的因果机制",
        "第五节以两个消亡案例（古埃及、拜占庭）和一个极端存续案例（犹太）验证外壳断裂的因果机制", "M5 结构概述纠错")
tr(P[60], "对比B：古埃及——宏观组织2.5+信仰2.5+外壳0=政权崩溃后消亡。",
        "对比B：古埃及——组织能力3.0（微观0.5/宏观2.5）+信仰2.5+外壳0=政权崩溃后消亡。", "M6 数字口径统一")
tr(P[67], "五个功能件中，信仰锚点和文化认同体系构成“核心”；血缘传递、组织能力和实践形式构成“外壳”。",
        "五个功能件中，信仰锚点与文化认同体系构成“核心”，回答“我们是谁”与“我们为何存在”；"
        "血缘身份传递、组织能力及其实践形式构成“外壳”，是核心承诺在具体历史条件下的承载形式；"
        "条件五（外壳调整）则是“在核心不变的前提下更换外壳”的制度化接口。", "M7 核心/外壳边界概念澄清")
tr(P[111], "但5a=0和5b=0使得这些努力无法突破锁死——", "但奠基期锁死的5a=0、5b=0使这些努力无从突破——", "M8 重复表述删改")
tr(P[113], "关于分析单位与组织形态的分类，详见补充材料D。",
        "拜占庭核心承诺被锁死的过程及其与古埃及的对照分析，详见补充材料C表C2。", "M9 脚注内容与正文不匹配修正")
tr(P[143], "完整的十三领域证据链（含兵制演变和经学演变）", "完整的十三领域证据链（含服饰、兵制和经学演变）", "M10 表注补正")
tr(P[151], "崇祯十六年（1643年），族人预储银两。",
        "早在崇祯十六年（1643年），族人即预储银两，为易代之际的宗族存续预作准备；", "M11 过程追踪时序澄清")
replace_para_tracked(P[168],
    "跨案例对照排除了冲击策略的混杂效应。如上节所示，犹太与古埃及面对的同为军事征服型冲击，"
    "且犹太所受的政治消灭更为彻底（耶路撒冷被夷为平地，圣殿被彻底摧毁，居民被驱逐流散），但两者结果相反——"
    "差异来自结构防御能力，而非冲击策略。中华案例则提供了“同一文明面对不同冲击策略”的内部对照——"
    "先后面对蒙古的军事征服与满清的正统性获取，冲击策略截然不同，但两次冲击后均成功维持了认同连续性。"
    "如果冲击策略是决定性变量，同一文明面对不同策略时应出现不同结果——但中华文明在两种策略下均存活。")
log.append("M12 §7.3与§7.2重复内容压缩")
tr(P[205], "没有“进步”——只需找到一个外壳≤1.0却在非极端冲击下成功重建了宏观政治组织的案例，框架即可被证伪。",
        "没有“进步”。而框架的可证伪性恰恰在于：只需找到一个外壳≤1.0却在非极端冲击下成功重建了宏观政治组织的案例，即可将其推翻。",
        "M13 语句衔接修正")
replace_para_tracked(P[175],
    "本文与既有理论的差异集中于两点：一是追问对象从“权力如何被行使与积累”转向“文明认同如何在权力消失后延续”；"
    "二是引入“奠基期写入的结构性参数”这一事前变量，使解释不再依赖危机中的事后表现。各理论的具体对照见表2及第二节。")
log.append("M14 §8.2扩写")

d.save(dst)
print("saved", dst)
for x in log: print(" -", x)

# ---------------- 补充材料 ----------------
src2 = "网络补充材料.docx"; dst2 = "/home/work/.openclaw/workspace/网络补充材料_修改稿.docx"
d2 = Document(src2)
Q = d2.paragraphs
log2 = []
def tr2(old, new, tag, style=None):
    p = find_para(d2, old, style)
    ok = replace_in_para(p, old, new)
    if not ok:
        replace_para_tracked(p, p.text.replace(old, new))
    log2.append(tag)

tr2("5e=0（从未提出过调整方案）", "5e=0（调整方案从未获得执行）", "S1 5e定义口径统一")

# S2 插入 A.1.3 小节标题
ref = find_para(d2, "关于观测可及性偏差的方法论声明")
insert_para_tracked(d2, ref, "A.1.3 条件三：文化认同体系（0—5.0）", "Heading 4", after=False)
log2.append("S2 补上缺失的A.1.3小节标题")

tr2("评级对照", "表A6d 条件三各案例评级对照", "S3 无编号表题补号")
tr2("表A8a（连接机制/关系遗产）结构类型分类", "表A10a（连接机制/关系遗产）结构类型分类", "S4 表A8a重复编号改正")
tr2("见补充材料C第C.4节", "见补充材料C第C.3节（表C2a及其说明）", "S5 交叉引用指向修正")
tr2("七、经学演变：从训诂学到考据学", "7.经学演变：从训诂学到考据学", "S6 编号体例统一")
tr2("7.跨领域模式一致性的理论意义", "8.跨领域模式一致性的理论意义", "S6 编号体例统一")
tr2("十二个领域的变革证据表明", "十三个领域的变革证据表明", "S7 领域计数与表C11一致")

# S11 表A7题注样式修正
pA7 = find_para(d2, "表A7 条件三评分标准（引擎B：仪式-身体层）")
pA7.style = d2.styles['表序与表题']
log2.append("S11 表A7题注样式修正")

# S8 表C5 题注样式修正 + 补建表格正文
cap = find_para(d2, "表C5 中华财政制度变革与外壳调整子维度的对应关系")
cap.style = d2.styles['表序与表题']
ref_tbl = d2.tables[51]  # 表C6，作为样式参照
tbl = d2.add_table(rows=1, cols=3)
try:
    tbl.style = ref_tbl.style
except Exception:
    tbl.style = 'Table Grid'
hdr = tbl.rows[0].cells
hdr[0].text = "外壳调整子维度"
hdr[1].text = "判定"
hdr[2].text = "千年财政史证据"
rows = [
    ("5c（调整的坚韧性）", "1.0",
     "从“初税亩”到“摊丁入亩”，跨越两千余年，每一个重大历史时期都有至少一次系统性财政改革，从未出现“放弃调整”的阶段"),
    ("5d（调整的方向多样性）", "1.0",
     "至少四次实质性转变：从人头税到土地税；从单纯征税到国家调控市场；从实物税到货币税；从人头税到财产税的最终完成"),
    ("5e（调整的完成度）", "0.5",
     "多数改革得到制度化执行，但部分重大改革执行中断或变形——王安石变法被废止、元朝纸币因滥发而崩溃"),
    ("核心承诺", "不变",
     "上述全部变革均未触及封建土地所有制的伦理基础、皇权专制的政治架构与儒家伦理的文化框架"),
]
for a, b, c in rows:
    cells = tbl.add_row().cells
    cells[0].text = a; cells[1].text = b; cells[2].text = c
cap._p.addnext(tbl._tbl)
log2.append("S8 表C5补建表格正文（依据上下文整理，需作者核对）")

# S9 表B7 题注（表格正文已存在但缺题注）
refb = find_para(d2, "1、不满足“文明”定义的判定理由")
insert_para_tracked(d2, refb, "表B7 哈扎尔帝国不满足“文明”定义的判定标准对照", "表序与表题", after=True)
log2.append("S9 补上缺失的表B7题注")

# S10 补上缺失的 D.2 小节标题
refd = find_para(d2, "表D2 横向嵌入型连接机制（中华案例）的运作逻辑")
insert_para_tracked(d2, refd, "D.2 三种连接机制的运作逻辑", "Heading 3", after=False)
log2.append("S10 补上缺失的D.2小节标题")
d2.save(dst2)
print("saved", dst2)
for x in log2: print(" -", x)

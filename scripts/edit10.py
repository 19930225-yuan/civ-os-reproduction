# -*- coding: utf-8 -*-
"""第九轮：补充材料E新增 E.15 Seshat 大样本映射（三期实测）"""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = "AI编辑"; DATE = "2026-10-17T07:20:00Z"
_c = [12000]
def nid():
    _c[0] += 1
    return str(_c[0])

def add_tracked_para(doc, text, style):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)._r
    parent = r.getparent(); idx = list(parent).index(r)
    parent.remove(r)
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), nid()); ins.set(qn('w:author'), AUTHOR); ins.set(qn('w:date'), DATE)
    ins.append(r); parent.insert(idx, ins)
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        st = pPr.find(qn('w:pStyle'))
        (st.addnext(rPr) if st is not None else pPr.insert(0, rPr))
    m = OxmlElement('w:ins')
    m.set(qn('w:id'), nid()); m.set(qn('w:author'), AUTHOR); m.set(qn('w:date'), DATE)
    rPr.append(m)
    return p

sup = Document("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")
for text, style in [
 ("E.15 Seshat 大样本映射：三期实测", "Heading 3"),
 ("数据与映射：Seshat 全球历史数据库公开快照（社会复杂度数据集 2017、道德宗教数据集 2021、轴心时代数据集 2018），覆盖 30 个历史自然地理区（NGA）、398 个政权。"
  "五功能件映射：条件一→神圣文本/宗教文献/专职祭司/宗教等级（辅以道德化超自然惩罚 MSP）；条件三→文字/书面记录/史学/文学/哲学；"
  "条件四→行政层级/全职官僚/职业军官/常备军/驿站；条件五→法典/法官/法庭/科举/考核（制度化调整通道的代理）；"
  "5a 的主代理→轴心数据集“统治者非神”。条件二（血缘）在 Seshat 无对应变量，为映射缺口。数据与分析文件：实测数据/seshat/。", "Normal"),
 ("（一）首期：政权存续时长——阴性结果与因变量纠偏。以政权存续时长为因变量（N=241），条件一/三/四与时长不相关，条件五代理为负相关（rho=−0.21，p=0.001），"
  "Logit P(存续≥200年) 全不显著（伪R²=0.02）。按本文框架这不是证伪而是因变量错置：政权本是可替换的外壳（中国王朝短命而文明长存），"
  "“政权长寿”并非预测对象；制度复杂政权更替频繁而文明延续，与“外壳可更换”隐喻相容。正确的因变量是崩溃后的文化保留。", "Normal"),
 ("（二）二期：转型配对（“壳碎核移”的大样本操作化）。NGA 内政权崩溃/更替前后配对 N=130（27 个 NGA），因变量为崩溃后文化复杂度变化（ΔC3，原始量纲 0—2）。"
  "三个结果：（1）描述性主发现——ΔC3 均值仅 +0.034，即文化复杂度在政权转型间近乎完全保留，“壳碎核移”是跨文明常态而非例外；"
  "（2）原始相关为负（C5_pre×ΔC3 rho=−0.28），但基线检查显示 C5_pre×C3_pre rho=+0.65，差值受均值回归污染；"
  "（3）控制基线后翻正：ΔC3 ~ C3_pre + C5_pre，C5_pre 系数 +0.117（p=0.013）——崩溃前制度化调整通道越强，崩溃后文化保留越好，"
  "为外壳断裂律的“壳碎核移”方向提供大样本初步支持（效应量中等偏弱）。", "Normal"),
 ("（三）三期：5a 主代理与 NGA 层。以“统治者非神”替换制度复杂度作 5a 主代理（可匹配子样本 N=76、8 个 NGA）：系数 +0.005（p=0.29），效应微弱且不显著——"
  "更贴近理论的 5a 代理在有限覆盖下未能复制 C5 的效应，如实报告。NGA 层文化保留率（ΣC3_post/ΣC3_pre）：多数 NGA 恰为 1.000（完全保留），"
  "瓦哈卡谷 1.67（转型后文化增长）、冰岛 1.20、上埃及 0.987（唯一下降者——恰为本文古埃及案例的大样本影子）、近畿 1.000（其“统治者非神”均值为 0，"
  "即天皇神权结构下文化仍完全保留，与互补律“组织/文化通道可补偿”相容）；NGA 层相关 rho=+0.27（p=0.51，N=8）不显著。", "Normal"),
 ("（四）限定与结论。三点限定：（1）映射是代理——Seshat 制度复杂度不完全等于 5a/5b“核心-外壳可分离性”，5a 主代理的覆盖又有限；"
  "（2）条件二映射缺口；（3）轴心数据集仅 10 个 NGA。结论：大样本层面最稳固的是描述性事实“政权转型间文化近乎完全保留”（壳碎核移为常态），"
  "其次是“制度化调整通道→崩溃后文化保留”的条件性效应（p=0.013）；5a 的独立效应待覆盖更好的数据检验。"
  "首期阴性结果与三期的不显著一并保留——选择性报告不可接受。", "Normal"),
]:
    add_tracked_para(sup, text, style)
sup.save("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx")
print("E.15 已写入")

def accept_all(src, dst):
    dd = Document(src)
    body = dd.element.body
    for ins in list(body.findall('.//' + qn('w:ins'))):
        parent = ins.getparent(); idx = list(parent).index(ins)
        if parent.tag == qn('w:rPr'):
            parent.remove(ins); continue
        for ch in list(ins):
            parent.insert(idx, ch); idx += 1
        parent.remove(ins)
    for dl in list(body.findall('.//' + qn('w:del'))):
        parent = dl.getparent()
        if parent.tag == qn('w:rPr'):
            pp = parent.getparent().getparent()
            has_text = any(t.text for t in pp.findall('.//' + qn('w:t')))
            if not has_text:
                pp.getparent().remove(pp)
            else:
                parent.remove(dl)
            continue
        parent.remove(dl)
    dd.save(dst)
accept_all("/home/work/.openclaw/workspace/投稿69_修改稿.docx", "/home/work/.openclaw/workspace/投稿69_终稿.docx")
accept_all("/home/work/.openclaw/workspace/网络补充材料_修改稿.docx", "/home/work/.openclaw/workspace/网络补充材料_终稿.docx")
print("终稿已更新")
